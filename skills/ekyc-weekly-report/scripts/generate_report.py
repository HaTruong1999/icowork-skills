#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generate the VNPT eKYC weekly HTML report from the weekly input .xlsx.

Usage:
    python3 generate_report.py <input.xlsx> <output.html>

Expects the xlsx to follow the layout of assets/input_template.xlsx:
  - Table "KR năm"        : header row has 'Mã KR' in some cell
  - Table "Hoạt động"     : header row has 'Hoạt động / tính năng'
  - Section "THÔNG TIN CHUNG"            : Trường | Giá trị rows
  - Section "CHỈ SỐ VẬN HÀNH TUẦN"       : metric rows
  - Section "ĐIỂM NỔI BẬT TRONG TUẦN"    : highlight rows

All sections are optional except the KR and Hoạt động tables — if a section
is missing, sensible placeholder content is used instead so the HTML is
always valid and ready for manual editing in the browser.
"""
import sys
import os
import re
import html
import openpyxl

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ASSETS_DIR = os.path.join(os.path.dirname(SCRIPT_DIR), "assets")

CARD_COLORS = ["blue", "teal", "purple", "amb", "red"]
ACTIVITY_STATUS_COLOR = {
    "đang triển khai": "blue",
    "đang vận hành": "teal",
    "hoàn thành": "teal",
    "đã xong": "teal",
    "trễ": "red",
    "chậm": "red",
    "cần hỗ trợ": "amb",
    "rủi ro": "amb",
    "tạm dừng": "amb",
}
HIGHLIGHT_TYPES = {
    "hợp đồng": ("contract", "💼"),
    "hop dong": ("contract", "💼"),
    "sự cố": ("incident", "🚨"),
    "su co": ("incident", "🚨"),
    "nâng cấp": ("milestone", "🚀"),
    "nang cap": ("milestone", "🚀"),
}
OVERALL_STATUS_COLOR = [
    (("rủi ro", "trễ", "chậm"), "red"),
    (("lưu ý", "cảnh báo"), "amb"),
    (("tốt", "ổn định", "on dinh", "tot"), "teal"),
]


def esc(v):
    if v is None:
        return ""
    return html.escape(str(v).strip())


def bold_markdown(text):
    """Convert **bold** markers to <b>bold</b>, escaping everything else."""
    text = "" if text is None else str(text)
    parts = re.split(r"\*\*(.+?)\*\*", text)
    out = []
    for i, p in enumerate(parts):
        if i % 2 == 1:
            out.append("<b>{}</b>".format(html.escape(p)))
        else:
            out.append(html.escape(p))
    return "".join(out).replace("\n", "<br>")


def lines_to_ul(text):
    """Turn a newline-separated cell (bullet lines starting with '-') into a <ul class="blist">."""
    if not text:
        return ""
    raw_lines = [l.strip() for l in str(text).split("\n") if l.strip()]
    if not raw_lines:
        return ""
    items = []
    for l in raw_lines:
        l = re.sub(r"^[-•+]\s*", "", l)
        items.append("<li contenteditable=\"false\">{}</li>".format(bold_markdown(l)))
    return "<ul class=\"blist\">{}</ul>".format("".join(items))


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("%", "")
    try:
        return float(s)
    except ValueError:
        return None


def pick_color(idx):
    return CARD_COLORS[idx % len(CARD_COLORS)]


def find_header_row(ws, must_contain):
    """Return the row index whose cell values contain `must_contain` (case-insens, stripped)."""
    for row in ws.iter_rows():
        vals = [str(c.value).strip().lower() if c.value is not None else "" for c in row]
        if must_contain.lower() in vals:
            return row[0].row
    return None


def read_table(ws, header_row):
    """Read a header row + subsequent data rows until a fully-empty row."""
    headers = []
    max_col = ws.max_column
    for c in range(1, max_col + 1):
        v = ws.cell(row=header_row, column=c).value
        headers.append(str(v).strip() if v is not None else "")
    # trim trailing empty headers
    while headers and headers[-1] == "":
        headers.pop()
    rows = []
    r = header_row + 1
    while r <= ws.max_row:
        vals = [ws.cell(row=r, column=c).value for c in range(1, len(headers) + 1)]
        if all(v is None or str(v).strip() == "" for v in vals):
            break
        row_dict = {}
        for h, v in zip(headers, vals):
            if h:
                row_dict[h] = v
        rows.append(row_dict)
        r += 1
    return rows


def read_key_value_section(ws, title_marker):
    """Read a 'Trường | Giá trị' style section that starts after a title row containing title_marker."""
    title_row = None
    for row in ws.iter_rows():
        if row[0].value and title_marker.lower() in str(row[0].value).strip().lower():
            title_row = row[0].row
            break
    if title_row is None:
        return {}
    header_row = title_row + 1
    rows = read_table(ws, header_row)
    out = {}
    for row in rows:
        key = None
        val = None
        for k, v in row.items():
            if k.lower().startswith("trường"):
                key = v
            elif k.lower().startswith("giá trị"):
                val = v
        if key:
            out[str(key).strip()] = val
    return out


def read_section_table(ws, title_marker):
    title_row = None
    for row in ws.iter_rows():
        if row[0].value and title_marker.lower() in str(row[0].value).strip().lower():
            title_row = row[0].row
            break
    if title_row is None:
        return []
    header_row = title_row + 1
    return read_table(ws, header_row)


# --------------------------------------------------------------------
# Section builders
# --------------------------------------------------------------------

def build_header(info, kr_rows):
    gd = esc(info.get("GĐ Sản phẩm") or "—")
    ky = esc(info.get("Kỳ báo cáo") or "—")
    status_text = str(info.get("Trạng thái chung tuần") or "Đang cập nhật").strip()
    color = "blue"
    low = status_text.lower()
    for keywords, c in OVERALL_STATUS_COLOR:
        if any(k in low for k in keywords):
            color = c
            break
    return """<header>
  <div class="row">
    <div>
      <h1 contenteditable="false">Báo cáo Tuần — VNPT eKYC</h1>
      <div class="sub" contenteditable="false">
        <span contenteditable="false">👤 <b>GĐ Sản phẩm:</b> <span class="val-txt" contenteditable="false">{gd}</span></span>
        <span contenteditable="false">•</span>
        <span contenteditable="false">📅 <b>Kỳ báo cáo:</b> <span class="val-txt" contenteditable="false">{ky}</span></span>
      </div>
    </div>
    <div class="status">
      <div style="font-size:11px;color:#cbd5e1;text-transform:uppercase;letter-spacing:.3px">Trạng thái chung tuần</div>
      <span class="pill {color} lg" id="overall-status" title="Click để sửa chữ. Nhấp đúp (Double-click) để đổi màu" contenteditable="false">{status}</span>
    </div>
  </div>
</header>""".format(gd=gd, ky=ky, color=color, status=esc(status_text))


def build_metric_cards(metric_rows):
    if not metric_rows:
        return "", 0
    cards = []
    for i, row in enumerate(metric_rows):
        name = esc(row.get("Tên chỉ số"))
        if not name:
            continue
        icon = str(row.get("Icon (tuỳ chọn)") or "📊").strip()
        color = str(row.get("Màu thẻ (tuỳ chọn)") or "").strip().lower()
        if color not in CARD_COLORS:
            color = pick_color(i)
        val = esc(row.get("Giá trị tuần này") or "—")
        prev = row.get("Giá trị tuần trước")
        pct = row.get("% thay đổi (tuỳ chọn)")
        trend = str(row.get("Xu hướng") or "").strip().lower()
        tag = esc(row.get("Ghi chú / tag") or "")

        if trend == "tăng-tốt":
            arrow, cls = "▲", "trend-up"
        elif trend == "giảm-tốt":
            arrow, cls = "▼", "trend-good-down"
        elif trend == "tăng-xấu":
            arrow, cls = "▲", "trend-down"
        elif trend == "giảm-xấu":
            arrow, cls = "▼", "trend-down"
        else:
            arrow, cls = "▬", "trend-neutral"

        sub_bits = []
        if pct not in (None, ""):
            pct_str = str(pct).strip()
            if not pct_str.startswith(("+", "-")) and cls in ("trend-up",):
                pct_str = "+" + pct_str
            sub_bits.append('<span class="{}">{} {}</span>'.format(cls, arrow, esc(pct_str)))
        else:
            sub_bits.append('<span class="{}">{}</span>'.format(cls, arrow if trend else "▬"))
        if prev not in (None, ""):
            sub_bits.append("vs tuần trước ({})".format(esc(prev)))
        sub_html = " ".join(sub_bits)

        cards.append("""    <div class="metric-card {color}">
      <div class="metric-title" contenteditable="false">{name} <span>{icon}</span></div>
      <div class="metric-value" contenteditable="false">{val}</div>
      <div class="metric-sub" contenteditable="false">{sub}</div>
      <span class="metric-tag" contenteditable="false">{tag}</span>
    </div>""".format(color=color, name=name.upper(), icon=icon, val=val, sub=sub_html, tag=tag))
    return "\n".join(cards), len(cards)


def build_highlights(highlight_rows):
    if not highlight_rows:
        return """      <div class="highlight-item milestone" contenteditable="false">
        📌 <b>Chưa có dữ liệu điểm nổi bật tuần này.</b> Thêm bằng nút "➕ Thêm điểm nổi bật" bên dưới.
      </div>"""
    items = []
    for row in highlight_rows:
        content = None
        type_val = None
        for k, v in row.items():
            kl = k.lower()
            if kl.startswith("loại"):
                type_val = v
            elif kl.startswith("nội dung"):
                content = v
        if not content:
            continue
        type_key = str(type_val or "").strip().lower()
        cls, icon = HIGHLIGHT_TYPES.get(type_key, ("milestone", "📌"))
        items.append('      <div class="highlight-item {cls}" contenteditable="false">\n        {icon} {content}\n      </div>'.format(
            cls=cls, icon=icon, content=bold_markdown(content)))
    return "\n".join(items) if items else build_highlights([])


def build_activities(activity_rows):
    if not activity_rows:
        return "", "", ""
    trs = []
    risk_boxes = []
    next_week_items = []
    for row in activity_rows:
        name = row.get("Hoạt động / tính năng")
        if not name:
            continue
        status = str(row.get("Trạng thái") or "").strip()
        status_l = status.lower()
        pill_color = "blue"
        for k, c in ACTIVITY_STATUS_COLOR.items():
            if k in status_l:
                pill_color = c
                break
        kr_tag = esc(row.get("Gắn KR / BAU / Khác") or "")
        done = row.get("Đã làm")
        upcoming = row.get("Sắp tới")
        blocker = row.get("Vướng")

        detail_html = lines_to_ul(done) if done else ""
        if not detail_html and done:
            detail_html = "<p contenteditable=\"false\">{}</p>".format(bold_markdown(done))

        trs.append("""        <tr>
          <td contenteditable="false">
            <b>{name}</b>
            <div style="font-size:11px;color:var(--muted);margin-top:2px">{tag}</div>
          </td>
          <td contenteditable="false"><span class="pill {color}" title="Click để sửa chữ. Nhấp đúp để đổi màu" contenteditable="false">{status}</span></td>
          <td contenteditable="false">
            {detail}
          </td>
        </tr>""".format(name=esc(name), tag=kr_tag, color=pill_color, status=esc(status) or "Đang triển khai",
                         detail=detail_html))

        if upcoming and str(upcoming).strip():
            next_week_items.append('<li contenteditable="false"><b>{name}:</b> {content}</li>'.format(
                name=esc(name), content=bold_markdown(upcoming)))

        if blocker and str(blocker).strip():
            risk_boxes.append("""    <div class="alert-box warn">
      <b contenteditable="false">🟡 {name}</b>
      <ul class="blist">
        <li contenteditable="false"><b contenteditable="false">Hiện trạng:</b> {content}</li>
      </ul>
      <div class="alert-action" contenteditable="false">
        👉 <b contenteditable="false">Đề xuất Sếp chỉ đạo:</b> <span contenteditable="false">[Điền đề xuất cụ thể]</span>
      </div>
    </div>""".format(name=esc(name), content=bold_markdown(blocker)))

    return "\n".join(trs), "\n".join(risk_boxes), "\n".join(next_week_items)


def kr_status_from_progress(progress):
    if progress is None:
        return "Có rủi ro", "amb"
    if progress >= 85:
        return "Đúng tiến độ", "teal"
    if progress >= 50:
        return "Có rủi ro", "amb"
    return "Trễ tiến độ", "red"


def compute_kr_progress(target, curr, prev):
    value = curr if curr not in (None, "") else prev
    t = num(target)
    v = num(value)
    if t is None or v is None:
        # qualitative KR (e.g. target "GTM"): treat v as a 0..1 ratio, or already a %
        if v is None:
            return None
        return min(100.0, max(0.0, v * 100 if v <= 1 else v))
    if t == 0:
        return None
    return min(100.0, max(0.0, (v / t) * 100))


def build_kr_section(kr_rows):
    if not kr_rows:
        summary = """  <div class="kr-stat-pill">
      <span contenteditable="false">Tổng số KR theo dõi</span>
    </div>"""
        return "0", "", "", ""
    cards = []
    table_rows = []
    counts = {"teal": 0, "amb": 0, "red": 0}
    for row in kr_rows:
        code = row.get("Mã KR")
        if not code:
            continue
        name = row.get("Key Result")
        unit = row.get("Đơn vị")
        target = row.get("Target 2026")
        prev = row.get("Tuần trước")
        curr = row.get("Tuần này")
        manual_status = str(row.get("Trạng thái (tuỳ chọn)") or "").strip()

        progress = compute_kr_progress(target, curr, prev)
        if manual_status in ("Đúng tiến độ", "Có rủi ro", "Trễ tiến độ"):
            status_text = manual_status
            color = {"Đúng tiến độ": "teal", "Có rủi ro": "amb", "Trễ tiến độ": "red"}[manual_status]
        else:
            status_text, color = kr_status_from_progress(progress)
        counts[color] += 1
        bar_width = 100 if progress is None else round(progress, 1)

        cards.append("""  <div class="kr-card">
    <div class="kr-top">
      <div>
        <span class="kr-code" contenteditable="false">{code}</span>
        <span class="kr-name" style="margin-left:6px" contenteditable="false">{name}</span>
      </div>
      <span class="pill {color}" title="Click để sửa chữ. Nhấp đúp để đổi màu" contenteditable="false">{status}</span>
    </div>
    <div class="kr-bar-track"><div class="kr-bar-fill {color}" style="width:{width}%"></div></div>
    <div class="kr-meta"><span contenteditable="false">Target: {target}</span><span contenteditable="false">Tuần trước: {prev}</span></div>
  </div>""".format(code=esc(code), name=esc(name), color=color, status=esc(status_text),
                    width=bar_width, target=esc(target) if target not in (None, "") else "—",
                    prev=esc(prev) if prev not in (None, "") else "-"))

        table_rows.append("""        <tr>
          <td contenteditable="false"><span class="kr-code" contenteditable="false">{code}</span></td>
          <td contenteditable="false">{name}</td>
          <td contenteditable="false">{unit}</td>
          <td contenteditable="false">{target}</td>
          <td contenteditable="false">{prev}</td>
          <td contenteditable="false">{curr}</td>
          <td contenteditable="false"><span class="pill {color}" title="Click để sửa chữ. Nhấp đúp để đổi màu" contenteditable="false">{status}</span></td>
        </tr>""".format(code=esc(code), name=esc(name), unit=esc(unit), target=esc(target),
                         prev=esc(prev) if prev not in (None, "") else "-",
                         curr=esc(curr) if curr not in (None, "") else "-",
                         color=color, status=esc(status_text)))

    total = len(cards)
    return str(total), "\n".join(cards), "\n".join(table_rows), counts


def build_next_week(items_html):
    if not items_html:
        return '<li contenteditable="false">Chưa có kế hoạch — thêm bằng nút "➕ Thêm mục kế hoạch" bên dưới.</li>'
    return items_html


def build_risks(risk_html):
    if not risk_html:
        return """    <div class="alert-box info">
      <b contenteditable="false">🟢 Không có vướng mắc nào được ghi nhận trong tuần này.</b>
    </div>"""
    return risk_html


def main():
    if len(sys.argv) != 3:
        print("Usage: python3 generate_report.py <input.xlsx> <output.html>")
        sys.exit(1)
    xlsx_path, out_path = sys.argv[1], sys.argv[2]

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    kr_header = find_header_row(ws, "Mã KR")
    kr_rows = read_table(ws, kr_header) if kr_header else []

    act_header = find_header_row(ws, "Hoạt động / tính năng")
    activity_rows = read_table(ws, act_header) if act_header else []

    info = read_key_value_section(ws, "THÔNG TIN CHUNG")
    metric_rows = read_section_table(ws, "CHỈ SỐ VẬN HÀNH TUẦN")
    highlight_rows = read_section_table(ws, "ĐIỂM NỔI BẬT")

    header_html = build_header(info, kr_rows)
    metric_cards_html, metric_count = build_metric_cards(metric_rows)
    highlights_html = build_highlights(highlight_rows)
    activities_html, risks_html, next_week_html = build_activities(activity_rows)
    kr_total, kr_cards_html, kr_table_rows_html, kr_counts = build_kr_section(kr_rows)

    if isinstance(kr_counts, dict):
        c_teal, c_amb, c_red = kr_counts["teal"], kr_counts["amb"], kr_counts["red"]
    else:
        c_teal = c_amb = c_red = 0

    main_html = """<main contenteditable="true">

<!-- ============================================================ -->
<!-- 1. CHỈ SỐ SỨC KHOẺ & VẬN HÀNH TUẦN (WEEK-ON-WEEK) -->
<!-- ============================================================ -->
<section id="sec-weekly-metrics">
  <div class="sec-header">
    <h2 contenteditable="false">⚡ 1. Chỉ số Sức khoẻ &amp; Vận hành Tuần (Week-on-Week)</h2>
    <span class="sec-tagline" contenteditable="false"><br></span>
  </div>

  <div class="metric-grid" id="metric-cards-container" data-count="{metric_count}">
{metric_cards}
  </div>

  <div class="tbl-actions" style="margin-top:6px">
    <button class="btn-tbl" id="btn-add-metric-card">➕ Thêm ô chỉ số</button>
  </div>

  <div class="highlights-container">
    <div class="highlights-title" contenteditable="false">📌 Điểm nổi bật trong tuần (Highlights &amp; Sự kiện trọng tâm)</div>
    <div id="highlights-list">
{highlights}
    </div>
    <div class="tbl-actions">
      <button class="btn-tbl" id="btn-add-highlight">➕ Thêm điểm nổi bật</button>
    </div>
  </div>
</section>

<!-- ============================================================ -->
<!-- 2. TIẾN ĐỘ HOẠT ĐỘNG TRỌNG TÂM TRONG TUẦN -->
<!-- ============================================================ -->
<section id="sec-activities">
  <div class="sec-header">
    <h2 contenteditable="false">🚀 2. Tiến độ Hoạt động Trọng tâm trong Tuần</h2>
    <span class="sec-tagline" contenteditable="false"><br></span>
  </div>

  <div class="tbl-scroll">
    <table id="tbl-activities">
      <thead>
        <tr>
          <th style="width:28%" contenteditable="false">Hoạt động / Tính năng</th>
          <th style="width:16%" contenteditable="false">Trạng thái</th>
          <th style="width:56%" contenteditable="false">Nội dung chi tiết</th>
        </tr>
      </thead>
      <tbody>
{activities}
      </tbody>
    </table>
  </div>

  <div class="tbl-actions">
    <button class="btn-tbl" id="btn-add-activity">➕ Thêm dòng công việc</button>
  </div>
</section>

<!-- ============================================================ -->
<!-- 3. VẤN ĐỀ, RỦI RO & ĐỀ XUẤT CẦN SẾP CHỈ ĐẠO (BLOCKERS) -->
<!-- ============================================================ -->
<section id="sec-risks">
  <div class="sec-header">
    <h2 contenteditable="false">⚠️ 3. Vấn đề, Rủi ro &amp; Đề xuất Hỗ trợ (Blockers &amp; Escalations)</h2>
    <span class="sec-tagline" contenteditable="false"><br></span>
  </div>

  <div id="risk-container">
{risks}
  </div>

  <div class="tbl-actions">
    <button class="btn-tbl" id="btn-add-risk">➕ Thêm vấn đề / Rủi ro mới</button>
  </div>
</section>

<!-- ============================================================ -->
<!-- 4. KẾ HOẠCH TRỌNG TÂM TUẦN TỚI -->
<!-- ============================================================ -->
<section id="sec-next-week">
  <div class="sec-header">
    <h2 contenteditable="false">🎯 4. Kế hoạch Trọng tâm Tuần tới</h2>
    <span class="sec-tagline" contenteditable="false"><br></span>
  </div>

  <div style="background:#fff;border:1px solid var(--line);border-radius:10px;padding:14px 18px;box-shadow:var(--shadow-sm)">
    <ul class="blist" id="list-next-week">
{next_week}
    </ul>
  </div>

  <div class="tbl-actions">
    <button class="btn-tbl" id="btn-add-plan">➕ Thêm mục kế hoạch</button>
  </div>
</section>

<!-- ============================================================ -->
<!-- 5. TỔNG QUAN & TIẾN ĐỘ TÍCH LUỸ KR NĂM (ANNUAL KR TRACKING) -->
<!-- ============================================================ -->
<section id="sec-annual-kr" style="margin-top:36px;border-top:2px dashed var(--line);padding-top:24px">
  <div class="sec-header" style="border-bottom:2px solid var(--navy)">
    <h2 contenteditable="false">📌 5. Tổng quan &amp; Tiến độ Tích luỹ KR Năm (Annual KR Tracking)</h2>
  </div>
  <div class="kr-summary-bar">
    <div class="kr-stat-pill">
      <b contenteditable="false">{kr_total}</b>
      <span contenteditable="false">Tổng số KR theo dõi</span>
    </div>
    <div class="kr-stat-pill teal">
      <b contenteditable="false">{c_teal}</b>
      <span contenteditable="false">Đúng tiến độ</span>
    </div>
    <div class="kr-stat-pill amb">
      <b contenteditable="false">{c_amb}</b>
      <span contenteditable="false">Có rủi ro</span>
    </div>
    <div class="kr-stat-pill red">
      <b contenteditable="false">{c_red}</b>
      <span contenteditable="false">Trễ tiến độ</span>
    </div>
  </div>

{kr_cards}

  <div style="margin-top:14px" class="tbl-scroll">
    <table id="tbl-kr-annual">
      <thead>
        <tr>
          <th style="width:10%" contenteditable="false">Mã KR</th>
          <th style="width:38%" contenteditable="false">Key Result (Mục tiêu chiến lược 2026)</th>
          <th style="width:10%" contenteditable="false">Đơn vị</th>
          <th style="width:12%" contenteditable="false">Target 2026</th>
          <th style="width:10%" contenteditable="false">Tuần trước</th>
          <th style="width:10%" contenteditable="false">Tuần này</th>
          <th style="width:10%" contenteditable="false">Trạng thái</th>
        </tr>
      </thead>
      <tbody>
{kr_table_rows}
      </tbody>
    </table>
  </div>
</section>

</main>""".format(
        metric_count=metric_count,
        metric_cards=metric_cards_html,
        highlights=highlights_html,
        activities=activities_html,
        risks=build_risks(risks_html),
        next_week=build_next_week(next_week_html),
        kr_total=kr_total, c_teal=c_teal, c_amb=c_amb, c_red=c_red,
        kr_cards=kr_cards_html,
        kr_table_rows=kr_table_rows_html,
    )

    with open(os.path.join(ASSETS_DIR, "part_top.html"), encoding="utf-8") as f:
        top = f.read()
    with open(os.path.join(ASSETS_DIR, "part_bottom.html"), encoding="utf-8") as f:
        bottom = f.read()

    full_html = top + header_html + "\n\n" + main_html + "\n\n" + bottom

    with open(out_path, "w", encoding="utf-8") as f:
        f.write(full_html)
    print("Wrote", out_path)


if __name__ == "__main__":
    main()
