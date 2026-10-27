#!/usr/bin/env python3
"""Trích xuất nội dung + ảnh trang từ tài liệu gốc (pptx, pdf, docx, html, xlsx, txt, md).

Thường được gọi qua `report.py extract`. Dùng riêng:
  python3 extract.py <file> <out_dir> [--dpi 60] [--no-images] [--no-convert]
  python3 extract.py <file.pdf> <out_dir> --render-pages 9,35 --dpi 110   # render nét vài trang

Kết quả trong <out_dir>:
  content.md      chữ theo từng slide/trang (## Slide N / ## Trang N), kèm ghi chú người nói của pptx
  metadata.json   loại file, số trang, thuộc tính (title, subject, author...), file anh em cùng tên
  pages/          ảnh từng trang (page-01.png ...), nếu render được
  sheets/         contact sheet 6 trang/ảnh (sheet-01.png ...) để xem nhanh toàn bộ
  hires/          ảnh nét các trang chỉ định bằng --render-pages

Backend PDF: ưu tiên PyMuPDF (fitz); nếu không có thì dùng poppler (pdftotext/pdftoppm/pdfinfo).
pptx/docx không có PDF đi kèm: đổi sang PDF bằng LibreOffice (soffice) nếu có, để lấy ảnh slide.
"""
import argparse
import json
import os
import platform
import re
import shutil
import subprocess
import sys
from html.parser import HTMLParser
from pathlib import Path

try:
    import fitz  # PyMuPDF
except Exception:
    fitz = None


# ---------------------------------------------------------------- tiện ích

def find_soffice():
    p = shutil.which("soffice") or shutil.which("libreoffice")
    if p:
        return p
    candidates = [
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    return next((c for c in candidates if os.path.exists(c)), None)


def run(cmd):
    return subprocess.run(cmd, capture_output=True, text=True, check=False)


def iso(dt):
    return dt.isoformat() if dt else None


# ---------------------------------------------------------------- PDF

def pdf_metadata(path):
    if fitz:
        doc = fitz.open(path)
        meta = {k: v for k, v in (doc.metadata or {}).items() if v}
        return doc.page_count, meta
    if shutil.which("pdfinfo"):
        out = run(["pdfinfo", str(path)]).stdout
        meta = {}
        for line in out.splitlines():
            if ":" in line:
                k, v = line.split(":", 1)
                meta[k.strip()] = v.strip()
        return int(meta.get("Pages", 0) or 0), meta
    return 0, {"warning": "không có PyMuPDF/pdfinfo"}


def pdf_pages_text(path):
    if fitz:
        doc = fitz.open(path)
        return [p.get_text("text", sort=True) for p in doc]
    if shutil.which("pdftotext"):
        out = run(["pdftotext", "-layout", str(path), "-"]).stdout
        pages = out.split("\f")
        if pages and not pages[-1].strip():
            pages = pages[:-1]
        return pages
    raise SystemExit("Không đọc được PDF: cần PyMuPDF (pip install PyMuPDF) hoặc poppler (pdftotext).")


def render_pdf(path, out_dir, dpi, pages=None):
    """Render trang PDF ra PNG. pages: list số trang (1-based) hoặc None = tất cả."""
    out_dir.mkdir(parents=True, exist_ok=True)
    written = []
    if fitz:
        doc = fitz.open(path)
        nums = pages or range(1, doc.page_count + 1)
        for n in nums:
            if 1 <= n <= doc.page_count:
                f = out_dir / f"page-{n:02d}.png"
                doc[n - 1].get_pixmap(dpi=dpi).save(f)
                written.append(f)
        return written
    if not shutil.which("pdftoppm"):
        print("! Không render được ảnh: thiếu PyMuPDF và pdftoppm", file=sys.stderr)
        return written
    tmp = out_dir / "_tmp"
    tmp.mkdir(exist_ok=True)
    ranges = [(n, n) for n in pages] if pages else [(None, None)]
    for a, b in ranges:
        cmd = ["pdftoppm", "-r", str(dpi), "-png"]
        if a:
            cmd += ["-f", str(a), "-l", str(b)]
        run(cmd + [str(path), str(tmp / "p")])
    for f in sorted(tmp.glob("p-*.png")):
        n = int(re.search(r"-(\d+)\.png$", f.name).group(1))
        dest = out_dir / f"page-{n:02d}.png"
        f.replace(dest)
        written.append(dest)
    shutil.rmtree(tmp, ignore_errors=True)
    return written


def contact_sheets(page_files, sheet_dir, per_sheet=6, cols=2, width=800):
    try:
        from PIL import Image, ImageDraw
    except Exception:
        print("! Thiếu Pillow: bỏ qua contact sheet", file=sys.stderr)
        return []
    sheet_dir.mkdir(parents=True, exist_ok=True)
    files = sorted(page_files)
    sheets = []
    for k in range(0, len(files), per_sheet):
        group = files[k:k + per_sheet]
        thumbs = []
        for f in group:
            im = Image.open(f).convert("RGB")
            h = int(im.height * width / im.width)
            thumbs.append((f, im.resize((width, h))))
        band = 22  # nhãn số trang nằm trên dải riêng, không che nội dung slide
        th = max(t.height for _, t in thumbs) + band
        rows = (len(thumbs) + cols - 1) // cols
        sheet = Image.new("RGB", (width * cols, th * rows), "white")
        draw = ImageDraw.Draw(sheet)
        for i, (f, t) in enumerate(thumbs):
            x, y = (i % cols) * width, (i // cols) * th
            n = re.search(r"(\d+)\.png$", f.name).group(1)
            draw.rectangle([x, y, x + width, y + band - 1], fill="black")
            draw.text((x + 6, y + 5), f"Trang/slide {int(n)}", fill="yellow")
            sheet.paste(t, (x, y + band))
        out = sheet_dir / f"sheet-{k // per_sheet + 1:02d}.png"
        sheet.save(out)
        sheets.append(out)
    return sheets


# ---------------------------------------------------------------- PPTX

def pptx_shape_texts(shapes, acc):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            pptx_shape_texts(sh.shapes, acc)
            continue
        top = sh.top or 0
        left = sh.left or 0
        if getattr(sh, "has_table", False) and sh.has_table:
            rows = []
            for r in sh.table.rows:
                rows.append(" | ".join(c.text.replace("\n", " ").strip() for c in r.cells))
            acc.append((top, left, "\n".join(rows)))
        elif getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t:
                acc.append((top, left, t))


def extract_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    cp = prs.core_properties
    props = {
        "title": cp.title, "subject": cp.subject, "author": cp.author,
        "last_modified_by": cp.last_modified_by, "keywords": cp.keywords,
        "created": iso(cp.created), "modified": iso(cp.modified),
    }
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        acc = []
        pptx_shape_texts(slide.shapes, acc)
        # đọc theo thứ tự trên → dưới, trái → phải (gom theo dải ~0.3 inch)
        acc.sort(key=lambda x: (round(x[0] / 274320), x[1]))
        body = "\n".join(t for _, _, t in acc)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        block = f"## Slide {i}\n\n{body}\n"
        if notes:
            block += "\n> Ghi chú người nói: " + notes.replace("\n", "\n> ") + "\n"
        parts.append(block)
    return len(prs.slides), props, parts


# ---------------------------------------------------------------- DOCX

def extract_docx(path):
    import docx
    d = docx.Document(path)
    cp = d.core_properties
    props = {
        "title": cp.title, "subject": cp.subject, "author": cp.author,
        "last_modified_by": cp.last_modified_by, "keywords": cp.keywords,
        "created": iso(cp.created), "modified": iso(cp.modified),
    }
    lines = []
    body = d.element.body
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(child, d)
            t = p.text.strip()
            if not t:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            m = re.search(r"(Heading|Tiêu đề)\s*(\d)", style)
            lines.append(("#" * (int(m.group(2)) + 1) + " " + t) if m else t)
        elif tag == "tbl":
            tbl = Table(child, d)
            for r in tbl.rows:
                cells = []
                for c in r.cells:
                    cells.append(c.text.replace("\n", " ").strip())
                lines.append("| " + " | ".join(cells) + " |")
            lines.append("")
    n_images = len(d.inline_shapes)
    props["inline_images"] = n_images
    return None, props, ["## Toàn văn\n\n" + "\n".join(lines) + "\n"]


# ---------------------------------------------------------------- HTML

class _HTMLText(HTMLParser):
    SKIP = {"script", "style", "noscript", "template", "svg"}
    BLOCK = {"p", "div", "li", "tr", "section", "article", "br", "h1", "h2", "h3",
             "h4", "h5", "h6", "td", "th", "header", "footer", "figcaption", "blockquote"}

    def __init__(self):
        super().__init__()
        self.out, self.skip, self.title, self.meta = [], 0, "", {}
        self._in_title = False
        self._heading = None

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag in self.SKIP:
            self.skip += 1
        if tag == "title":
            self._in_title = True
        if tag == "meta" and a.get("name") and a.get("content"):
            self.meta[a["name"]] = a["content"]
        if tag in self.BLOCK:
            self.out.append("\n")
        if re.fullmatch(r"h[1-6]", tag):
            self._heading = int(tag[1])
            self.out.append("#" * (self._heading + 1) + " ")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skip:
            self.skip -= 1
        if tag == "title":
            self._in_title = False
        if tag in ("td", "th"):
            self.out.append(" | ")
        if tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if self._in_title:
            self.title += data
            return
        if self.skip:
            return
        t = re.sub(r"\s+", " ", data)
        if t.strip():
            self.out.append(t)


def extract_html(path):
    raw = Path(path).read_text(encoding="utf-8", errors="replace")
    p = _HTMLText()
    p.feed(raw)
    text = re.sub(r"\n\s*\n+", "\n\n", "".join(p.out)).strip()
    props = {"title": p.title.strip(), **{f"meta:{k}": v for k, v in p.meta.items()}}
    props["has_inline_script"] = "<script" in raw.lower()
    return None, props, ["## Toàn văn\n\n" + text + "\n"]


def compare_versions(parts, n_slides, pdf):
    """So PDF cùng tên với pptx/docx: số trang + tỷ lệ đoạn chữ của từng slide tìm thấy trong PDF."""
    try:
        n_pdf, _ = pdf_metadata(pdf)
        pdf_text = re.sub(r"\s+", "", "".join(pdf_pages_text(pdf)))
    except SystemExit:
        return "Không đọc được PDF để so phiên bản."
    samples = []
    for block in parts:
        body = block.split("\n", 2)[-1].split("> Ghi chú người nói:")[0]
        lines = [l.strip() for l in body.splitlines() if len(l.strip()) >= 12]
        if lines:
            samples.append(re.sub(r"\s+", "", lines[0])[:30])
    found = sum(1 for s in samples if s in pdf_text)
    ratio = found / len(samples) if samples else 0
    same_pages = (n_slides is None) or (n_pdf == n_slides)
    verdict = "KHỚP" if same_pages and ratio >= 0.8 else "CÓ THỂ KHÁC PHIÊN BẢN"
    return (f"So phiên bản: {verdict} (PDF {n_pdf} trang / gốc {n_slides if n_slides else '?'}; "
            f"{found}/{len(samples)} đoạn mẫu có trong PDF).")


# ---------------------------------------------------------------- XLSX / TXT

def extract_xlsx(path):
    try:
        import openpyxl
    except Exception:
        raise SystemExit("Đọc .xlsx cần openpyxl: python3 -m pip install openpyxl")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in r]
            if any(cells):
                rows.append("| " + " | ".join(cells) + " |")
        parts.append(f"## Sheet {ws.title}\n\n" + "\n".join(rows) + "\n")
    return len(parts), {"sheets": len(parts)}, parts


def extract_text(path):
    t = Path(path).read_text(encoding="utf-8", errors="replace")
    return None, {}, ["## Toàn văn\n\n" + t.strip() + "\n"]


# ---------------------------------------------------------------- main

def to_pdf_with_soffice(src, out_dir):
    soffice = find_soffice()
    if not soffice:
        return None
    r = run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(src)])
    pdf = out_dir / (Path(src).stem + ".pdf")
    if pdf.exists():
        return pdf
    print("! LibreOffice không đổi được sang PDF:", r.stderr[-300:], file=sys.stderr)
    return None


def extract_file(src, out, dpi=60, no_images=False, no_convert=False):
    """Trích một file vào thư mục out. Trả về metadata (dict); chữ nằm ở out/content.md."""
    src = Path(src).expanduser().resolve()
    out = Path(out).expanduser().resolve()
    out.mkdir(parents=True, exist_ok=True)
    ext = src.suffix.lower()
    warnings = []
    if ext == ".pdf":
        n, props = pdf_metadata(src)
        texts = pdf_pages_text(src)
        parts = [f"## Trang {i}\n\n{t.rstrip()}\n" for i, t in enumerate(texts, 1)]
        n = n or len(texts)
        kind = "pdf"
        if sum(len(t.strip()) for t in texts) < 30 * max(1, len(texts)):
            warnings.append("PDF gần như không có chữ (có thể là ảnh scan): xem ảnh trang trong pages/ hoặc OCR trước.")
    elif ext in (".pptx", ".pptm"):
        n, props, parts = extract_pptx(src)
        kind = "pptx"
    elif ext == ".docx":
        n, props, parts = extract_docx(src)
        kind = "docx"
    elif ext in (".html", ".htm"):
        n, props, parts = extract_html(src)
        kind = "html"
    elif ext in (".xlsx", ".xlsm"):
        n, props, parts = extract_xlsx(src)
        kind = "xlsx"
    elif ext in (".txt", ".md", ".csv"):
        n, props, parts = extract_text(src)
        kind = "text"
    elif ext in (".ppt", ".doc", ".xls"):
        raise SystemExit(f"Định dạng cũ ({ext}): đổi sang pptx/docx/xlsx hoặc PDF trước "
                         "(LibreOffice: soffice --headless --convert-to pdf <file>).")
    else:
        raise SystemExit(f"Chưa hỗ trợ định dạng {ext}")

    page_files, sheets, pdf_for_images = [], [], None
    if not no_images:
        if kind == "pdf":
            pdf_for_images = src
        elif kind in ("pptx", "docx"):
            sib = src.with_suffix(".pdf")
            if sib.exists():
                pdf_for_images = sib
                warnings.append(f"Dùng {sib.name} (PDF cùng tên) để lấy ảnh trang. " + compare_versions(parts, n, sib))
            elif not no_convert:
                pdf_for_images = to_pdf_with_soffice(src, out)
        if pdf_for_images:
            page_files = render_pdf(pdf_for_images, out / "pages", dpi)
            sheets = contact_sheets(page_files, out / "sheets")

    meta = {
        "source": str(src), "file_name": src.name, "type": kind,
        "size_bytes": src.stat().st_size, "pages": n, "properties": props,
        "page_images": len(page_files), "contact_sheets": [str(s) for s in sheets],
        "warnings": warnings,
    }
    (out / "metadata.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    header = [f"# {src.name}", "", f"- Loại: {kind} · Số trang/slide: {n}"]
    for k, v in props.items():
        if v not in (None, ""):
            header.append(f"- {k}: {v}")
    (out / "content.md").write_text("\n".join(header) + "\n\n" + "\n".join(parts), encoding="utf-8")
    return meta


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("out_dir")
    ap.add_argument("--dpi", type=int, default=60)
    ap.add_argument("--no-images", action="store_true", help="chỉ lấy chữ + metadata")
    ap.add_argument("--no-convert", action="store_true", help="không gọi LibreOffice")
    ap.add_argument("--render-pages", help="chỉ render các trang này (vd 9,35) vào hires/ rồi thoát")
    args = ap.parse_args()
    src = Path(args.file).expanduser().resolve()
    out = Path(args.out_dir).expanduser().resolve()
    if args.render_pages:
        pdf = src if src.suffix.lower() == ".pdf" else src.with_suffix(".pdf")
        if not pdf.exists():
            pdf = to_pdf_with_soffice(src, out)
        if not pdf:
            raise SystemExit("Không có PDF để render.")
        pages = [int(x) for x in args.render_pages.split(",") if x.strip()]
        print("\n".join(str(f) for f in render_pdf(pdf, out / "hires", args.dpi, pages)))
        return
    meta = extract_file(src, out, args.dpi, args.no_images, args.no_convert)
    print(f"Đã trích: {out / 'content.md'}")
    for w in meta["warnings"]:
        print("! " + w)


if __name__ == "__main__":
    main()
