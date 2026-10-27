#!/usr/bin/env python3
"""report.json -> Word (.docx) kiểu đơn giản: đủ chữ, không trang trí.

Mỗi slide = một mục (Heading 1); mỗi khối = tiêu đề phụ (Heading 2) + đoạn văn / gạch đầu dòng / bảng
lưới mặc định. Chỉ giữ chữ đậm từ **...**. Cần python-docx.
Thường được gọi qua `report.py docx`. Dùng riêng:
  python3 to_docx.py <report.json> <out.docx>
"""
import json
import re
import sys
from pathlib import Path


def add_runs(par, text):
    """Thêm chữ vào đoạn, **...** thành chữ đậm."""
    for k, piece in enumerate(re.split(r"\*\*(.+?)\*\*", str(text))):
        if piece:
            par.add_run(piece).bold = (k % 2 == 1)
    return par


def para(doc, text, style=None):
    return add_runs(doc.add_paragraph(style=style), text)


def bullet(doc, text):
    return para(doc, text, "List Bullet")


def table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for c, h in enumerate(header):
        t.rows[0].cells[c].text = ""
        add_runs(t.rows[0].cells[c].paragraphs[0], f"**{h}**" if h else "")
    for r in rows:
        cells = t.add_row().cells
        for c, v in enumerate(r):
            add_runs(cells[c].paragraphs[0], v)
    doc.add_paragraph()


def heading(doc, text, level):
    h = doc.add_heading(level=level)
    add_runs(h, re.sub(r"\*\*", "", str(text)))


def block(doc, b):
    t = b["type"]
    if t == "stack":
        for x in b["blocks"]:
            block(doc, x)
        return
    if b.get("title") and t not in ("category", "panel"):
        heading(doc, b["title"], 2)
    if t == "stat":
        bullet(doc, f'**{b["value"]}** – {b["label"]}')
    elif t == "split":
        for p in b["parts"]:
            line = f'**{p["value"]} {p["label"]}**'
            if p.get("desc"):
                line += f': {p["desc"]}'
            bullet(doc, line)
        if b.get("note"):
            para(doc, b["note"])
    elif t == "numlist":
        for it in b["items"]:
            bullet(doc, f'**{it["value"]}** – {it["text"]}')
    elif t == "category":
        head = b["title"] + (f' – {b["total"]}' if b.get("total") else "")
        heading(doc, head, 2)
        if b.get("sub"):
            para(doc, b["sub"])
        table(doc, ["Nhóm", "Số lượng"], [[r["label"], r["value"]] for r in b["rows"]])
    elif t == "progress":
        para(doc, f'{b.get("label", "Tiến độ")}: **{b["value"]}**')
        if b.get("caption"):
            para(doc, b["caption"])
    elif t == "panel":
        heading(doc, b["title"], 2)
        if b.get("text"):
            para(doc, b["text"])
        for x in b.get("bullets") or []:
            bullet(doc, x)
    elif t == "status_table":
        table(doc, ["Hạng mục", "Trạng thái", "Nội dung"], [[r["item"], r["status"], r["desc"]] for r in b["rows"]])
    elif t == "status_list":
        for it in b["items"]:
            bullet(doc, f'**{it["status"]}:** {it["text"]}')
    elif t == "steps":
        for k, x in enumerate(b["items"], 1):
            para(doc, f"{k}. {x}")
    elif t == "text":
        para(doc, b["text"])
    elif t == "table":
        table(doc, b["columns"], b["rows"])
    elif t == "bars":
        for it in b["items"]:
            bullet(doc, f'{it["label"]}: **{it["value"]}**')


def build_docx(rep, out_path):
    import docx
    doc = docx.Document()
    meta = rep.get("meta", {})
    doc.core_properties.title = meta.get("title", "")
    doc.core_properties.author = meta.get("author", "")
    for i, s in enumerate(rep["slides"], 1):
        typ = s.get("type", "content")
        if typ == "cover":
            if s.get("kicker"):
                para(doc, s["kicker"])
            doc.add_heading(s["title"], 0)
            for k in ("subtitle", "date"):
                if s.get(k):
                    para(doc, s[k])
            if s.get("agenda"):
                para(doc, "**Nội dung:**")
                for a in s["agenda"]:
                    para(doc, f'{a.get("num", "")}. {a["text"]}'.lstrip(". "))
            continue
        if s.get("kicker"):
            para(doc, s["kicker"])
        heading(doc, s["title"], 1)
        if s.get("lead") or s.get("subtitle"):
            para(doc, s.get("lead") or s.get("subtitle"))
        for row in s.get("rows") or []:
            for b in row["blocks"]:
                block(doc, b)
        if s.get("next"):
            para(doc, f'**Tiếp theo:** {s["next"]}')
        if s.get("note"):
            n = s["note"]
            para(doc, f'**{n["label"]}:** {n["text"]}' if isinstance(n, dict) and n.get("label")
                 else (n["text"] if isinstance(n, dict) else n))
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    if len(sys.argv) < 3:
        raise SystemExit(__doc__)
    rep = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    print("Word:", build_docx(rep, sys.argv[2]))
